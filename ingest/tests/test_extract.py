"""Extraction without LibreOffice: label payloads, and the decks that fail before rendering."""

import zipfile

import pytest
from pptx import Presentation

from app.detect import LabelBox
from app.errors import UserFacingError
from app.extract import NO_SLIDES, NOT_A_DECK, diagram_payload, extract_deck
from app.geometry import Box


def test_labels_become_fractions_of_the_crop_in_reading_order():
    # 200 x 100 points of slide, cut out at zoom 2 as 400 x 200 pixels.
    frame = Box(100.0, 50.0, 300.0, 150.0)
    labels = [
        LabelBox("Lower", Box(150.0, 120.0, 200.0, 130.0)),
        LabelBox("Prefrontal\ncortex", Box(120.0, 60.0, 180.0, 70.0), 91.6),
    ]

    payload = diagram_payload(labels, frame, (400, 200))

    assert payload is not None
    assert payload.model_dump(mode="json") == {
        "labels": [
            {"text": "Prefrontal cortex", "x": 0.1, "y": 0.1, "w": 0.3, "h": 0.1, "confidence": 92},
            {"text": "Lower", "x": 0.25, "y": 0.7, "w": 0.25, "h": 0.1},
        ],
        "image_width": 400,
        "image_height": 200,
    }


def test_labels_are_clipped_to_the_crop_and_dropped_outside_it():
    frame = Box(0.0, 0.0, 100.0, 100.0)
    edge = LabelBox("Edge", Box(90.0, 10.0, 120.0, 20.0))
    outside = LabelBox("Outside", Box(150.0, 10.0, 160.0, 20.0))

    payload = diagram_payload([edge, outside], frame, (100, 100))

    assert payload is not None
    assert [(label.text, label.x, label.w) for label in payload.labels] == [("Edge", 0.9, 0.1)]
    assert diagram_payload([outside], frame, (100, 100)) is None


def test_a_renamed_text_file_fails_legibly(tmp_path):
    source = tmp_path / "lecture-04.pptx"
    source.write_text("Week 4 notes: the cerebrum has four lobes.")

    with pytest.raises(UserFacingError) as error:
        extract_deck(source, tmp_path)
    assert str(error.value) == NOT_A_DECK


def test_a_zip_that_isnt_a_deck_fails_legibly(tmp_path):
    source = tmp_path / "lecture-04.pptx"
    with zipfile.ZipFile(source, "w") as archive:
        archive.writestr("notes.txt", "not a presentation")

    with pytest.raises(UserFacingError) as error:
        extract_deck(source, tmp_path)
    assert str(error.value) == NOT_A_DECK


def test_a_deck_of_hidden_slides_fails_legibly(tmp_path):
    presentation = Presentation()
    hidden = presentation.slides.add_slide(presentation.slide_layouts[6])
    hidden.element.set("show", "0")
    source = tmp_path / "hidden.pptx"
    presentation.save(str(source))

    with pytest.raises(UserFacingError) as error:
        extract_deck(source, tmp_path)
    assert str(error.value) == NO_SLIDES
