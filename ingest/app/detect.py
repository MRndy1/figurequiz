"""Tier 1 labeled-diagram detection for PPTX (CLAUDE.md, "the core algorithm").

Reads shape geometry straight out of the slide XML: a picture with two or more short text
labels on or beside it is a labeled diagram. No computer vision, and no rendering — the
rendered page is only needed later, to snap label boxes onto the drawn text (see render.py).
Tables come out of the same pass, cell by cell, for table questions.

Everything here works in points, in slide coordinate space.
"""

from __future__ import annotations

from dataclasses import dataclass, field

from pptx.enum.shapes import PP_PLACEHOLDER
from pptx.oxml.ns import qn
from pptx.shapes.group import GroupShape
from pptx.shapes.picture import Picture as PptxPicture

from app.geometry import Box, emu_to_pt

# Placeholders that are page furniture, never diagram labels.
FURNITURE_PLACEHOLDERS = frozenset(
    {
        PP_PLACEHOLDER.TITLE,
        PP_PLACEHOLDER.CENTER_TITLE,
        PP_PLACEHOLDER.SUBTITLE,
        PP_PLACEHOLDER.FOOTER,
        PP_PLACEHOLDER.DATE,
        PP_PLACEHOLDER.SLIDE_NUMBER,
    }
)


@dataclass(frozen=True)
class DetectConfig:
    """Thresholds worth tuning against real decks (see `app.cli inspect`)."""

    max_label_words: int = 5
    min_labels_per_diagram: int = 2
    #: How far outside a picture a label may sit, as a fraction of slide width. Labels on
    #: leader lines sit well clear of the image, so this is deliberately generous.
    label_margin_frac: float = 0.06
    #: A picture must cover at least this fraction of the slide to be worth OCR when it has
    #: too few labels of its own — its labels may be baked into the pixels.
    min_ocr_picture_area_frac: float = 0.05
    #: Padding added around a diagram crop, as a fraction of slide width.
    crop_pad_frac: float = 0.02
    #: Without a title placeholder, the topmost text box is the title if it has at most this
    #: many words...
    max_title_words: int = 10
    #: ...and starts within this top fraction of the slide.
    title_zone_frac: float = 0.4

    # PDFs only (app.pdf). A deck keeps its master's pictures out of the slide's shapes, but a
    # PDF draws them on every page.
    #: Pictures smaller than this fraction of the page are logos and icons, not figures.
    min_picture_area_frac: float = 0.02
    #: A picture this wide or this tall, as a fraction of the page, is a background, banner or
    #: side bar, not a figure.
    page_furniture_frac: float = 0.9
    #: A page with fewer characters of text than this (CLAUDE.md, OCR routing)...
    scan_max_text_chars: int = 20
    #: ...whose pictures cover at least this fraction of it is a scan: OCR reads all of it.
    scan_min_picture_frac: float = 0.5

    # Tables, from a deck's table shapes or a PDF's ruled tables. Past these sizes a table is
    # a page of data rather than something to learn cell by cell, and too wide for a phone.
    max_table_rows: int = 20
    max_table_cols: int = 8


@dataclass(frozen=True)
class Transform:
    """Maps a group's child coordinate space onto the slide.

    PowerPoint lets a group be resized without touching its children: the children keep
    their own coordinate space (`chOff`/`chExt`) and the group scales it into place. Ignoring
    this puts every grouped label in the wrong spot on resized groups.
    """

    off_x: float = 0.0
    off_y: float = 0.0
    scale_x: float = 1.0
    scale_y: float = 1.0
    child_off_x: float = 0.0
    child_off_y: float = 0.0

    def apply(self, x: float, y: float) -> tuple[float, float]:
        return (
            self.off_x + (x - self.child_off_x) * self.scale_x,
            self.off_y + (y - self.child_off_y) * self.scale_y,
        )


@dataclass(frozen=True)
class LabelBox:
    text: str
    box: Box
    #: Tesseract's 0-100 confidence, on Tier 2 labels only.
    confidence: float | None = None


@dataclass(frozen=True)
class PictureBox:
    box: Box


@dataclass(frozen=True)
class Diagram:
    picture: PictureBox
    labels: tuple[LabelBox, ...]

    def crop(self, slide: Box, config: DetectConfig, keep_clear: Box | None = None) -> Box:
        """The region to cut out: the picture, its labels, and a little breathing room.

        `keep_clear` is usually the slide title. Padding above the topmost label often reaches
        into it, and half a heading sliced across the top of a question looks broken, so the
        crop is pulled back below it — but only when the whole of `keep_clear` sits above the
        content, so this can never clip a label.
        """
        region = self.picture.box
        for label in self.labels:
            region = region.union(label.box)

        pad = slide.width * config.crop_pad_frac
        box = region.expanded(pad).clipped_to(slide)

        if keep_clear is not None and box.y0 < keep_clear.y1:
            # Push the top down past the title, but never past the topmost label, so this
            # trims the heading without ever clipping something the question needs.
            box = Box(box.x0, min(max(box.y0, keep_clear.y1), region.y0), box.x1, box.y1)
        return box


@dataclass(frozen=True)
class TableGrid:
    """A table's cells as text, every row as long as the first. A merged cell's text is in its
    first cell, and the cells it covers are empty."""

    rows: tuple[tuple[str, ...], ...]
    #: 1 when the first row is headings, else 0.
    header_rows: int


#: Fewer cells than this to blank is not a question (web/lib/generation/validate.ts, BLANKS).
MIN_TABLE_BODY_CELLS = 2


def table_grid(rows: list[list[str]], header_rows: int, config: DetectConfig) -> TableGrid | None:
    """`rows` as a TableGrid, each cell's whitespace collapsed, or None when the table is too
    small, too big or too empty to quiz on. Rows are padded to the longest with empty cells."""
    cells = [[" ".join(cell.split()) for cell in row] for row in rows]
    width = max((len(row) for row in cells), default=0)
    if not (2 <= len(cells) <= config.max_table_rows and 2 <= width <= config.max_table_cols):
        return None
    body = sum(1 for row in cells[header_rows:] for cell in row if cell)
    if body < MIN_TABLE_BODY_CELLS:
        return None
    return TableGrid(tuple(tuple(row + [""] * (width - len(row))) for row in cells), header_rows)


@dataclass(frozen=True)
class SlideDetection:
    #: 1-based position among *visible* slides, which is also the PDF page number.
    index: int
    title: str | None
    #: Every other text frame and table on the slide, in the slide's own (z-)order, which is
    #: PowerPoint's reading order. One paragraph per line.
    text: str
    diagrams: tuple[Diagram, ...] = field(default_factory=tuple)
    ocr_candidates: tuple[PictureBox, ...] = field(default_factory=tuple)
    #: Where the title sits, so crops can avoid slicing through it.
    title_box: Box | None = None
    #: The slide's tables worth quizzing on, in the slide's own order.
    tables: tuple[TableGrid, ...] = field(default_factory=tuple)


def is_hidden(slide) -> bool:  # noqa: ANN001 - pptx Slide has no public type export
    """PowerPoint marks a skipped slide `show="0"`, and LibreOffice omits it from the PDF."""
    return slide.element.get("show") == "0"


def visible_slides(presentation) -> list:  # noqa: ANN001
    return [slide for slide in presentation.slides if not is_hidden(slide)]


def slide_box(presentation) -> Box:  # noqa: ANN001
    """The slide rectangle in points."""
    width, height = presentation.slide_width, presentation.slide_height
    if width is None or height is None:
        raise ValueError("presentation has no slide size")
    return Box(0.0, 0.0, emu_to_pt(width), emu_to_pt(height))


def _group_transform(group: GroupShape, parent: Transform) -> Transform:
    """Compose `parent` with the group's own child-space transform."""
    xfrm = group._element.find(qn("p:grpSpPr"))
    xfrm = None if xfrm is None else xfrm.find(qn("a:xfrm"))
    if xfrm is None:
        return parent

    off, ext = xfrm.find(qn("a:off")), xfrm.find(qn("a:ext"))
    child_off, child_ext = xfrm.find(qn("a:chOff")), xfrm.find(qn("a:chExt"))
    if off is None or ext is None or child_off is None or child_ext is None:
        return parent

    child_cx, child_cy = float(child_ext.get("cx")), float(child_ext.get("cy"))
    if child_cx == 0 or child_cy == 0:
        return parent

    origin_x, origin_y = parent.apply(float(off.get("x")), float(off.get("y")))
    return Transform(
        off_x=origin_x,
        off_y=origin_y,
        scale_x=float(ext.get("cx")) * parent.scale_x / child_cx,
        scale_y=float(ext.get("cy")) * parent.scale_y / child_cy,
        child_off_x=float(child_off.get("x")),
        child_off_y=float(child_off.get("y")),
    )


def _shape_box(shape, transform: Transform) -> Box | None:  # noqa: ANN001
    if None in (shape.left, shape.top, shape.width, shape.height):
        return None
    x0, y0 = transform.apply(shape.left, shape.top)
    x1, y1 = transform.apply(shape.left + shape.width, shape.top + shape.height)
    return Box.from_corners(emu_to_pt(x0), emu_to_pt(y0), emu_to_pt(x1), emu_to_pt(y1))


def _placeholder_type(shape):  # noqa: ANN001, ANN201
    if not shape.is_placeholder:
        return None
    try:
        return shape.placeholder_format.type
    except (AttributeError, ValueError):
        return None


def _is_label(shape, text: str, config: DetectConfig) -> bool:  # noqa: ANN001
    if not text or len(text.split()) > config.max_label_words:
        return False
    return _placeholder_type(shape) not in FURNITURE_PLACEHOLDERS


def _text_of(shape) -> str:  # noqa: ANN001
    """A text frame's text, one paragraph or line break per line."""
    return shape.text_frame.text.replace("\v", "\n").strip()


def _table_text(shape) -> str:  # noqa: ANN001
    rows = ("\t".join(cell.text for cell in row.cells) for row in shape.table.rows)
    return "\n".join(rows)


def _table_grid(shape, config: DetectConfig) -> TableGrid | None:  # noqa: ANN001
    """A table shape's cells. PowerPoint's "Header Row" option marks the first row headings."""
    table = shape.table
    rows = [["" if cell.is_spanned else cell.text for cell in row.cells] for row in table.rows]
    return table_grid(rows, 1 if table.first_row else 0, config)


def _find_title(slide, bounds: Box, config: DetectConfig):  # noqa: ANN001, ANN202
    """The shape that titles the slide: its title placeholder, else the topmost text box.

    Real decks often title a slide with a plain text box, so `shapes.title` is None or an
    empty placeholder, which renders as nothing (docs/prototype-questions.md, §4.6). Left
    unfound, that heading counts as a label of the picture under it — a false labeled diagram —
    and crops slice through it. The topmost text box only counts when it is short and near the
    top: a longer one is a heading or prose that can't be a label anyway, and settling for the
    next box down would steal a real label.
    """
    placeholder = slide.shapes.title
    if placeholder is not None and _text_of(placeholder):
        return placeholder

    topmost, topmost_box = None, None
    for shape in slide.shapes:  # top level only: titles are never grouped
        if not shape.has_text_frame or not _text_of(shape):
            continue
        box = _shape_box(shape, Transform())
        if box is not None and (topmost_box is None or box.y0 < topmost_box.y0):
            topmost, topmost_box = shape, box

    if topmost is None or len(_text_of(topmost).split()) > config.max_title_words:
        return None
    if topmost_box.y0 > bounds.y0 + bounds.height * config.title_zone_frac:
        return None
    return topmost


def _walk(
    shapes,  # noqa: ANN001
    transform: Transform,
    config: DetectConfig,
    pictures: list[PictureBox],
    labels: list[LabelBox],
    texts: list[str],
    tables: list[TableGrid],
    title_id: int | None,
) -> None:
    """Collect pictures, label candidates, text and tables, descending into groups.

    The title shape (`title_id`) is skipped: it is the slide's title, never a label or body
    text.
    """
    for shape in shapes:
        # `shape_type` raises on shapes python-pptx cannot classify, so never touch it.
        if isinstance(shape, GroupShape):
            inner = _group_transform(shape, transform)
            _walk(shape.shapes, inner, config, pictures, labels, texts, tables, title_id)
            continue

        if isinstance(shape, PptxPicture):
            box = _shape_box(shape, transform)
            if box is not None and box.area > 0:
                pictures.append(PictureBox(box))
            continue

        if shape.has_text_frame:
            text = _text_of(shape)
            if not text or shape.shape_id == title_id:
                continue
            texts.append(text)
            box = _shape_box(shape, transform)
            if box is not None and _is_label(shape, text, config):
                labels.append(LabelBox(text, box))
            continue

        if getattr(shape, "has_table", False):
            texts.append(_table_text(shape))
            grid = _table_grid(shape, config)
            if grid is not None:
                tables.append(grid)


def build_diagrams(
    pictures: list[PictureBox],
    labels: list[LabelBox],
    slide: Box,
    config: DetectConfig,
) -> tuple[tuple[Diagram, ...], tuple[PictureBox, ...]]:
    """Attach each label to its nearest picture, then split diagrams from OCR candidates.

    The same rule for a deck's slides and a PDF's pages (app.pdf).
    """
    margin = slide.width * config.label_margin_frac
    attached: dict[int, list[LabelBox]] = {i: [] for i in range(len(pictures))}

    for label in labels:
        nearest, nearest_gap = None, None
        for index, picture in enumerate(pictures):
            gap = picture.box.gap_to(label.box)
            if gap <= margin and (nearest_gap is None or gap < nearest_gap):
                nearest, nearest_gap = index, gap
        if nearest is not None:
            attached[nearest].append(label)

    diagrams: list[Diagram] = []
    ocr_candidates: list[PictureBox] = []
    min_ocr_area = slide.area * config.min_ocr_picture_area_frac

    for index, picture in enumerate(pictures):
        picture_labels = attached[index]
        if len(picture_labels) >= config.min_labels_per_diagram:
            diagrams.append(Diagram(picture, tuple(picture_labels)))
        elif picture.box.area >= min_ocr_area:
            # A big picture with no labels of its own: they may be baked into the pixels.
            ocr_candidates.append(picture)

    return tuple(diagrams), tuple(ocr_candidates)


def detect_slide(
    slide,  # noqa: ANN001
    index: int,
    slide_bounds: Box,
    config: DetectConfig | None = None,
) -> SlideDetection:
    """Find labeled diagrams on one slide. `index` is its 1-based visible position."""
    config = config or DetectConfig()
    pictures: list[PictureBox] = []
    labels: list[LabelBox] = []
    texts: list[str] = []
    tables: list[TableGrid] = []

    title_shape = _find_title(slide, slide_bounds, config)
    title_id = title_shape.shape_id if title_shape is not None else None
    _walk(slide.shapes, Transform(), config, pictures, labels, texts, tables, title_id)
    diagrams, ocr_candidates = build_diagrams(pictures, labels, slide_bounds, config)

    # A wrapped title reads as one line.
    title = " ".join(_text_of(title_shape).split()) if title_shape is not None else None
    title_box = _shape_box(title_shape, Transform()) if title_shape is not None else None

    return SlideDetection(
        index=index,
        title=title or None,
        text="\n".join(texts),
        diagrams=diagrams,
        ocr_candidates=ocr_candidates,
        title_box=title_box,
        tables=tuple(tables),
    )


def detect_presentation(
    presentation,  # noqa: ANN001
    config: DetectConfig | None = None,
) -> list[SlideDetection]:
    """Detect across every visible slide, numbered to match the rendered PDF's pages."""
    bounds = slide_box(presentation)
    return [
        detect_slide(slide, index, bounds, config)
        for index, slide in enumerate(visible_slides(presentation), start=1)
    ]
