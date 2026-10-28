"""Builds a synthetic lecture deck so detection can be exercised without real slides.

Real decks are copyrighted and stay out of git (`ingest/samples/`), so tests and the first
demo run against this instead.
"""

from __future__ import annotations

from pathlib import Path

from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.util import Inches, Pt

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)


def _draw_cell_diagram(path: Path, size: tuple[int, int] = (1200, 900)) -> Path:
    """A crude unlabeled 'cell' — the kind of figure a lecturer pastes onto a slide."""
    image = Image.new("RGB", size, (255, 255, 255))
    draw = ImageDraw.Draw(image)
    w, h = size

    for (x0, y0, x1, y1), fill, outline, width in (
        ((0.08, 0.10, 0.92, 0.90), (214, 234, 248), (41, 128, 185), 8),
        ((0.38, 0.36, 0.62, 0.64), (174, 214, 241), (31, 97, 141), 6),
        ((0.20, 0.22, 0.32, 0.34), (249, 231, 159), (183, 149, 11), 5),
        ((0.68, 0.62, 0.82, 0.74), (212, 239, 223), (20, 143, 119), 5),
    ):
        draw.ellipse((w * x0, h * y0, w * x1, h * y1), fill=fill, outline=outline, width=width)
    draw.line((w * 0.30, h * 0.70, w * 0.45, h * 0.78), fill=(125, 60, 152), width=6)
    image.save(path)
    return path


def build_sample_deck(path: Path) -> Path:
    """Write a four-slide deck: a labeled diagram, a bare figure, a grouped-label diagram,
    and a hidden slide (which LibreOffice leaves out of the PDF)."""
    path.parent.mkdir(parents=True, exist_ok=True)
    figure = _draw_cell_diagram(path.parent / "sample-figure.png")

    prs = Presentation()
    prs.slide_width, prs.slide_height = SLIDE_W, SLIDE_H
    title_only, blank = prs.slide_layouts[5], prs.slide_layouts[6]

    # 1. Picture with four labels around it — the case the hero feature is built for.
    slide = prs.slides.add_slide(title_only)
    slide.shapes.title.text = "Animal cell structure"
    slide.shapes.add_picture(
        str(figure), Inches(4.2), Inches(1.4), width=Inches(5.0), height=Inches(3.75)
    )
    labels = {
        "Nucleus": (Inches(6.1), Inches(2.9)),
        "Cell membrane": (Inches(3.0), Inches(1.5)),
        "Mitochondrion": (Inches(8.4), Inches(4.2)),
        "Cytoplasm": (Inches(3.1), Inches(4.4)),
    }
    for text, (left, top) in labels.items():
        box = slide.shapes.add_textbox(left, top, Inches(1.7), Inches(0.35))
        box.text_frame.text = text
        box.text_frame.paragraphs[0].runs[0].font.size = Pt(14)

    # 2. A figure with no separate labels: a Tier 2 OCR candidate, not a Tier 1 diagram.
    slide = prs.slides.add_slide(title_only)
    slide.shapes.title.text = "Scanned figure from the textbook"
    slide.shapes.add_picture(
        str(figure), Inches(3.6), Inches(1.6), width=Inches(6.0), height=Inches(4.5)
    )

    # 3. Labels inside a group, which need the group transform to land correctly.
    slide = prs.slides.add_slide(blank)
    slide.shapes.add_picture(
        str(figure), Inches(1.2), Inches(1.2), width=Inches(4.5), height=Inches(3.4)
    )
    group = slide.shapes.add_group_shape()
    for text, (left, top) in {
        "Vacuole": (Inches(0.6), Inches(1.3)),
        "Ribosome": (Inches(5.5), Inches(3.6)),
    }.items():
        box = group.shapes.add_textbox(left, top, Inches(1.6), Inches(0.35))
        box.text_frame.text = text
        box.text_frame.paragraphs[0].runs[0].font.size = Pt(14)

    # 4. Hidden slide: rendering must skip it so slide numbers keep matching PDF pages.
    hidden = prs.slides.add_slide(title_only)
    hidden.shapes.title.text = "Cut for time"
    hidden.element.set("show", "0")

    prs.save(str(path))
    return path


def _shrink_group_child_space(group, factor: float = 0.5) -> None:  # noqa: ANN001
    """Halve a group's child extent so its children render at 1/factor scale.

    python-pptx always writes an identity group transform, so a generated deck never
    exercises the `chOff`/`chExt` maths. Tests call this to create the resized case.
    """
    from pptx.oxml.ns import qn

    xfrm = group._element.find(qn("p:grpSpPr")).find(qn("a:xfrm"))
    child_ext = xfrm.find(qn("a:chExt"))
    child_ext.set("cx", str(int(int(child_ext.get("cx")) * factor)))
    child_ext.set("cy", str(int(int(child_ext.get("cy")) * factor)))


__all__ = ["_shrink_group_child_space", "build_sample_deck"]
