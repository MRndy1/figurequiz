"""Reading an upload into what the database stores: page renders, page text, labeled diagrams.

One pass over a PPTX or a PDF (CLAUDE.md, Ingestion pipeline). For a deck, python-pptx reads
the slide XML for text and Tier 1 diagrams, and LibreOffice renders the visible slides; a PDF
already is pages, and `app.pdf` reads their text layer and pictures. Either way OCR reads the
labels baked into pictures that have too few labels of their own (Tier 2), and each diagram is
cut out of its page with its label boxes as 0-1 fractions of the cut-out — the exact image the
blank job draws on and the student sees. A scanned PDF page has its text read by OCR too.
Tables come out cell by cell, from the slide XML or the PDF's ruled tables.

Nothing here knows about users, documents or Storage: the worker (`app.jobs`) names and stores
what this returns, and `app.cli blank` writes it to a folder.
"""

from __future__ import annotations

import zipfile
from collections.abc import Callable, Iterable
from dataclasses import dataclass
from pathlib import Path

import pymupdf
from lxml import etree
from pptx import Presentation
from pptx.exc import PythonPptxError

from app.contracts import (
    COORDINATE_PLACES,
    DiagramLabel,
    LabeledDiagramPayload,
    TablePayload,
    TextBlockPayload,
)
from app.detect import (
    DetectConfig,
    Diagram,
    LabelBox,
    TableGrid,
    detect_presentation,
    slide_box,
    visible_slides,
)
from app.errors import UserFacingError
from app.geometry import Box
from app.ocr import OcrConfig, ocr_page_text, ocr_picture
from app.pdf import detect_page, page_box
from app.render import (
    check_page_count,
    convert_to_pdf,
    crop_image,
    render_page,
    snap_label_box,
    zoom_for,
)

#: Long edge of a page render. Plenty to read a slide on any screen; diagram crops go up to
#: `render.MAX_CROP_PX` because their blanks have to land on small text.
PAGE_RENDER_PX = 1600
#: Pages read from one PDF. Each is rendered, stored and perhaps OCR'd, and a textbook can run
#: to a thousand pages, which would hold the worker for an hour and fill Storage.
MAX_PDF_PAGES = 100

NOT_A_DECK = (
    "This file isn't a PowerPoint deck we can open. Save it as .pptx, without a password, "
    "from PowerPoint, Keynote or Google Slides, and upload that copy."
)
NO_SLIDES = "This deck has no visible slides to read."
NOT_A_PDF = "This file isn't a PDF we can open. Save or export it as a PDF again, and upload that."
PDF_PASSWORD = "This PDF is password-protected. Save a copy without the password, and upload that."
NO_PAGES = "This PDF has no pages to read."
UNKNOWN_TYPE = "This kind of file can't be read. Upload a PowerPoint deck (.pptx) or a PDF."

_WHOLE_IMAGE = Box(0.0, 0.0, 1.0, 1.0)


@dataclass(frozen=True)
class DiagramCrop:
    """One labeled diagram, cut out of its page."""

    page: int
    #: 1-based among its page's diagrams: Tier 1 first, then Tier 2.
    order: int
    #: 1: labels from the file's own text (slide XML, PDF text layer). 2: labels OCR read off
    #: the pixels.
    tier: int
    #: The unblanked crop, a PNG in the work folder.
    image_path: Path
    payload: LabeledDiagramPayload


@dataclass(frozen=True)
class PageText:
    payload: TextBlockPayload
    #: 1: the file's own text. 2: OCR read it off a scanned page.
    tier: int


@dataclass(frozen=True)
class TableFound:
    """One table, read from the file's own structure: never OCR'd."""

    page: int
    #: 1-based among its page's tables, in the page's own order.
    order: int
    payload: TablePayload


@dataclass(frozen=True)
class Extraction:
    page_count: int
    #: Full page renders, PNGs in the work folder: `pages[0]` is page 1.
    pages: tuple[Path, ...]
    #: Each page's text, for the pages that have any.
    texts: dict[int, PageText]
    diagrams: tuple[DiagramCrop, ...]
    tables: tuple[TableFound, ...] = ()


#: A diagram found on a page: its tier, its picture, and its labels' best-known boxes.
_Found = tuple[int, Diagram, list[LabelBox]]


def extract_upload(
    source: Path,
    work_dir: Path,
    config: DetectConfig | None = None,
    ocr_config: OcrConfig | None = None,
) -> Extraction:
    """Read the upload at `source` as the kind of file its extension says it is: an upload's
    object key ends in its type (web/lib/contracts.ts, `objectKeys.upload`)."""
    read = READERS.get(source.suffix.lower())
    if read is None:
        raise UserFacingError(UNKNOWN_TYPE)
    return read(source, work_dir, config, ocr_config)


def extract_deck(
    source: Path,
    work_dir: Path,
    config: DetectConfig | None = None,
    ocr_config: OcrConfig | None = None,
) -> Extraction:
    """Read the deck at `source`, writing `pages/NNN.png` and `diagrams/NNN-K.png` (and
    LibreOffice's PDF) into `work_dir`.

    Raises `UserFacingError` when the file isn't a deck that can be read, or its slides can't
    be rendered.
    """
    config = config or DetectConfig()
    ocr_config = ocr_config or OcrConfig()

    presentation = _open_deck(source)
    slide_count = len(visible_slides(presentation))
    if slide_count == 0:
        raise UserFacingError(NO_SLIDES)
    detections = detect_presentation(presentation, config)
    slide_bounds = slide_box(presentation)
    page_zoom = zoom_for(slide_bounds, max_px=PAGE_RENDER_PX)

    pages_dir, diagrams_dir = _work_dirs(work_dir)
    pages: list[Path] = []
    texts: dict[int, PageText] = {}
    diagrams: list[DiagramCrop] = []
    tables: list[TableFound] = []

    with pymupdf.open(convert_to_pdf(source, work_dir)) as document:
        check_page_count(document, slide_count)

        for detection in detections:
            number = detection.index
            page = document[number - 1]
            pages.append(_save_page(page, number, page_zoom, pages_dir))

            if detection.title or detection.text:
                payload = TextBlockPayload(title=detection.title, text=detection.text)
                texts[number] = PageText(payload, tier=1)
            tables.extend(_tables(number, detection.tables))

            # Tier 1 boxes come from the XML and get snapped onto the drawn text. Tier 2 boxes
            # are already where OCR found the ink.
            found: list[_Found] = [
                (1, diagram, [label for label, _snapped in snapped_labels(page, diagram)])
                for diagram in detection.diagrams
            ]
            for candidate in detection.ocr_candidates:
                read = ocr_picture(page, candidate, slide_bounds, config, ocr_config)
                if read is not None:
                    found.append((2, read, list(read.labels)))

            diagrams.extend(
                _cut_out(
                    page, number, found, slide_bounds, config, detection.title_box, diagrams_dir
                )
            )

    return Extraction(slide_count, tuple(pages), texts, tuple(diagrams), tuple(tables))


def extract_pdf(
    source: Path,
    work_dir: Path,
    config: DetectConfig | None = None,
    ocr_config: OcrConfig | None = None,
) -> Extraction:
    """Read the PDF at `source`, writing `pages/NNN.png` and `diagrams/NNN-K.png` into
    `work_dir`.

    Raises `UserFacingError` when the file isn't a PDF that can be read, or has too many pages.
    """
    config = config or DetectConfig()
    ocr_config = ocr_config or OcrConfig()

    pages_dir, diagrams_dir = _work_dirs(work_dir)
    pages: list[Path] = []
    texts: dict[int, PageText] = {}
    diagrams: list[DiagramCrop] = []
    tables: list[TableFound] = []

    with _open_pdf(source) as document:
        for number, page in enumerate(document, start=1):
            bounds = page_box(page)
            pages.append(
                _save_page(page, number, zoom_for(bounds, max_px=PAGE_RENDER_PX), pages_dir)
            )
            detection = detect_page(page, config)

            if detection.scanned:
                text = ocr_page_text(page, bounds, ocr_config)
                if text:
                    texts[number] = PageText(TextBlockPayload(title=None, text=text), tier=2)
            elif detection.title or detection.text:
                payload = TextBlockPayload(title=detection.title, text=detection.text)
                texts[number] = PageText(payload, tier=1)
            tables.extend(_tables(number, detection.tables))

            # Text-layer boxes are the drawn text already, so Tier 1 needs no snapping.
            found: list[_Found] = [
                (1, diagram, list(diagram.labels)) for diagram in detection.diagrams
            ]
            for candidate in detection.ocr_candidates:
                read = ocr_picture(page, candidate, bounds, config, ocr_config)
                if read is None:
                    continue
                labels = [label for label in read.labels if not detection.has_text_at(label.box)]
                if len(labels) >= config.min_labels_per_diagram:
                    found.append((2, Diagram(read.picture, tuple(labels)), labels))

            diagrams.extend(
                _cut_out(page, number, found, bounds, config, detection.title_box, diagrams_dir)
            )
        page_count = document.page_count

    return Extraction(page_count, tuple(pages), texts, tuple(diagrams), tuple(tables))


#: How each kind of upload is read, by its file extension.
READERS: dict[str, Callable[..., Extraction]] = {".pptx": extract_deck, ".pdf": extract_pdf}


def snapped_labels(page: pymupdf.Page, diagram: Diagram) -> list[tuple[LabelBox, bool]]:
    """Each Tier 1 label with its best-known box, and whether that box was snapped to the drawn
    text (True) or is the XML box (False)."""
    out = []
    for label in diagram.labels:
        snapped = snap_label_box(page, label.text, label.box)
        out.append((LabelBox(label.text, snapped or label.box), snapped is not None))
    return out


def diagram_payload(
    labels: Iterable[LabelBox], frame: Box, size: tuple[int, int]
) -> LabeledDiagramPayload | None:
    """`labels`, in `frame`'s units, as the stored payload of the image cut from `frame`.

    Boxes become 0-1 fractions of that image, clipped to it, and labels go in reading order:
    their index is what an answer key's `label_index` points at. A label with nothing left
    inside the image is dropped. None when no label is left.
    """
    entries = []
    for label in sorted(labels, key=lambda label: (label.box.y0, label.box.x0)):
        inside = label.box.normalized_within(frame).intersection(_WHOLE_IMAGE)
        if inside is None:
            continue
        w, h = round(inside.width, COORDINATE_PLACES), round(inside.height, COORDINATE_PLACES)
        if w <= 0 or h <= 0:
            continue
        entries.append(
            DiagramLabel(
                # A label wrapped onto two lines is still one name.
                text=" ".join(label.text.split()),
                x=round(inside.x0, COORDINATE_PLACES),
                y=round(inside.y0, COORDINATE_PLACES),
                w=w,
                h=h,
                confidence=None if label.confidence is None else round(label.confidence),
            )
        )
    if not entries:
        return None
    width, height = size
    return LabeledDiagramPayload(labels=entries, image_width=width, image_height=height)


def _tables(number: int, grids: Iterable[TableGrid]) -> list[TableFound]:
    return [
        TableFound(
            number,
            order,
            TablePayload(rows=[list(row) for row in grid.rows], header_rows=grid.header_rows),
        )
        for order, grid in enumerate(grids, start=1)
    ]


def _work_dirs(work_dir: Path) -> tuple[Path, Path]:
    pages_dir, diagrams_dir = work_dir / "pages", work_dir / "diagrams"
    pages_dir.mkdir(parents=True, exist_ok=True)
    diagrams_dir.mkdir(parents=True, exist_ok=True)
    return pages_dir, diagrams_dir


def _save_page(page: pymupdf.Page, number: int, zoom: float, pages_dir: Path) -> Path:
    path = pages_dir / f"{number:03d}.png"
    render_page(page, zoom).save(path)
    return path


def _cut_out(
    page: pymupdf.Page,
    number: int,
    found: list[_Found],
    bounds: Box,
    config: DetectConfig,
    title_box: Box | None,
    diagrams_dir: Path,
) -> list[DiagramCrop]:
    """Crop each diagram found on page `number` out of it, numbered in the order found."""
    crops: list[DiagramCrop] = []
    for tier, diagram, labels in found:
        crop = diagram.crop(bounds, config, title_box)
        zoom = zoom_for(crop)
        image = crop_image(render_page(page, zoom), crop, zoom)
        payload = diagram_payload(labels, crop, image.size)
        if payload is None or len(payload.labels) < config.min_labels_per_diagram:
            continue

        order = len(crops) + 1
        image_path = diagrams_dir / f"{number:03d}-{order}.png"
        image.save(image_path)
        crops.append(DiagramCrop(number, order, tier, image_path, payload))
    return crops


def _open_deck(source: Path):  # noqa: ANN202 - pptx exports no Presentation type
    """Open the deck, or say legibly why it isn't one: a renamed file, an old .ppt, a
    password-protected deck (an encrypted container, not a zip) or a damaged one."""
    try:
        return Presentation(str(source))
    except (
        PythonPptxError,
        zipfile.BadZipFile,
        KeyError,
        ValueError,
        etree.XMLSyntaxError,
    ) as error:
        raise UserFacingError(NOT_A_DECK) from error


def _open_pdf(source: Path) -> pymupdf.Document:
    """Open the PDF, or say legibly why it can't be read. PyMuPDF repairs what damage it can;
    a PDF that only restricts printing or copying opens like any other."""
    try:
        document = pymupdf.open(source, filetype="pdf")
    except pymupdf.FileDataError as error:  # not a PDF at all, or an empty file
        raise UserFacingError(NOT_A_PDF) from error

    problem = None
    if document.needs_pass:
        problem = PDF_PASSWORD
    elif document.page_count == 0:
        problem = NO_PAGES
    elif document.page_count > MAX_PDF_PAGES:
        problem = (
            f"This PDF has {document.page_count} pages, and up to {MAX_PDF_PAGES} can be read. "
            "Split it into smaller PDFs, one per lecture or chapter, and upload those."
        )
    if problem is not None:
        document.close()
        raise UserFacingError(problem)
    return document
